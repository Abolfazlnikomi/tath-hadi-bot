from datetime import datetime
from io import BytesIO
from zipfile import ZipFile

import jdatetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from zoneinfo import ZoneInfo


HEADERS = {
    "students": ["نام و نام خانوادگی", "کد ملی", "شماره تماس"],
    "attendance": ["کد ملی", "نام و نام خانوادگی", "وضعیت (حاضر / غایب)"],
    "discounts": ["کد ملی", "نام و نام خانوادگی", "کد تخفیف", "توضیحات/عنوان دوره"],
}
DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


class ExcelError(ValueError):
    pass


def normalize(value):
    return str(value if value is not None else "").strip().translate(DIGITS)


def national_id(value):
    result = normalize(value)
    if len(result) != 10 or not result.isascii() or not result.isdigit():
        raise ExcelError("کد ملی باید دقیقاً ۱۰ رقم باشد؛ ستون را به صورت Text ذخیره کنید.")
    return result


def today():
    current = datetime.now(ZoneInfo("Asia/Tehran")).date()
    return current, jdatetime.date.fromgregorian(date=current).strftime("%Y/%m/%d")


def workbook(kind, rows=(), classroom_id=None, session_date=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "داده‌ها"
    sheet.sheet_view.rightToLeft = True
    sheet.append(HEADERS[kind])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245B78")
    for row in rows:
        sheet.append([str(value) if value is not None else "" for value in row])
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.data_type = "s"
            cell.number_format = "@"
    for column in "ABCD"[:len(HEADERS[kind])]:
        sheet.column_dimensions[column].width = 30
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    if kind == "attendance":
        from openpyxl.worksheet.datavalidation import DataValidation

        validation = DataValidation(type="list", formula1='"حاضر,غایب"')
        sheet.add_data_validation(validation)
        validation.add(f"C2:C{max(2, sheet.max_row)}")
    metadata = book.create_sheet("_meta")
    metadata.append(["kind", kind])
    metadata.append(["classroom_id", str(classroom_id or "")])
    metadata.append(["session_date", session_date or ""])
    metadata.sheet_state = "hidden"
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def parse(content, kind, classroom_id, max_rows=10000):
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                raise ExcelError("حجم بازشده فایل اکسل بیش از حد مجاز است.")
        book = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except ExcelError:
        raise
    except Exception as error:
        raise ExcelError("فایل اکسل معتبر با پسوند xlsx ارسال کنید.") from error
    try:
        if "_meta" not in book.sheetnames:
            raise ExcelError("از فایل نمونه دریافتی همین ربات استفاده کنید؛ شیت اطلاعات حذف شده است.")
        metadata = dict(book["_meta"].iter_rows(values_only=True))
        if metadata.get("kind") != kind or str(metadata.get("classroom_id")) != str(classroom_id):
            raise ExcelError("این فایل مربوط به کلاس یا عملیات انتخاب‌شده نیست.")
        session_date = metadata.get("session_date")
        if kind == "attendance":
            try:
                parsed = datetime.strptime(session_date, "%Y-%m-%d").date()
            except (ValueError, TypeError) as error:
                raise ExcelError("تاریخ فایل حضور و غیاب معتبر نیست.") from error
            if parsed > today()[0]:
                raise ExcelError("ثبت حضور و غیاب برای آینده مجاز نیست.")
        sheet = book["داده‌ها"]
        iterator = sheet.iter_rows()
        if [cell.value for cell in next(iterator)] != HEADERS[kind]:
            raise ExcelError("عنوان ستون‌ها را تغییر ندهید.")
        records, seen = [], set()
        for index, cells in enumerate(iterator, 2):
            if index > max_rows + 1:
                raise ExcelError(f"حداکثر {max_rows} ردیف مجاز است.")
            if not any(cell.value is not None for cell in cells):
                continue
            if any(cell.data_type == "f" for cell in cells):
                raise ExcelError(f"ردیف {index}: استفاده از فرمول مجاز نیست.")
            values = [normalize(cell.value) for cell in cells]
            try:
                if kind == "students":
                    name, code, phone = values
                    code = national_id(code)
                    if not name or len(name) > 200:
                        raise ExcelError("نام الزامی و حداکثر ۲۰۰ کاراکتر است.")
                    if not phone.isascii() or not phone.isdigit() or not 10 <= len(phone) <= 15:
                        raise ExcelError("شماره تماس باید ۱۰ تا ۱۵ رقم باشد؛ صفر ابتدایی را حفظ کنید.")
                    record = {"national_id": code, "name": name, "phone": phone}
                    key = code
                else:
                    code = national_id(values[0])
                    if kind == "attendance":
                        if values[2] not in ("حاضر", "غایب"):
                            raise ExcelError("وضعیت فقط حاضر یا غایب است؛ وضعیت خالی مجاز نیست.")
                        record = {"national_id": code, "name": values[1], "status": values[2]}
                        key = code
                    else:
                        if not values[2] or len(values[2]) > 200 or len(values[3]) > 2000:
                            raise ExcelError("کد تخفیف الزامی (حداکثر ۲۰۰) و توضیحات حداکثر ۲۰۰۰ کاراکتر است.")
                        record = {"national_id": code, "name": values[1], "code": values[2], "description": values[3]}
                        key = (code, values[2])
                if key in seen:
                    raise ExcelError("رکورد تکراری در فایل وجود دارد.")
                seen.add(key)
                records.append(record)
            except (ExcelError, ValueError) as error:
                raise ExcelError(f"ردیف {index}: {error}") from error
        if not records:
            raise ExcelError("فایل هیچ ردیف داده‌ای ندارد.")
        return {"kind": kind, "classroom_id": classroom_id, "session_date": session_date, "rows": records}
    except ExcelError:
        raise
    except Exception as error:
        raise ExcelError("ساختار فایل معتبر نیست؛ از فایل نمونه ربات استفاده کنید.") from error
    finally:
        book.close()
