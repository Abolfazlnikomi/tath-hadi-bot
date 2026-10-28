import os
import unittest
from datetime import timedelta
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select

from schoolbot import excel
from schoolbot.api import BaleError
from schoolbot.config import Config
from schoolbot.models import (
    Attendance, Base, Classroom, Conversation, Discount, Enrollment, Inbox, Outbox, Student, utcnow,
)
from schoolbot.runtime import Runner
from schoolbot.service import apply_import


class FakeAPI:
    def __init__(self):
        self.calls = []
        self.content = b""
        self.error = None

    def call(self, method, payload=None):
        if self.error:
            raise self.error
        self.calls.append((method, payload))
        return True

    def download(self, _document):
        return self.content

    def close(self):
        pass


class BotTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine(os.getenv("TEST_DATABASE_URL", "sqlite+pysqlite:///:memory:"))
        Base.metadata.drop_all(self.engine)
        self.api = FakeAPI()
        self.config = Config("test-token", "unused")
        self.runner = Runner(self.config, self.api, self.engine)
        self.runner.initialize()
        self.sessions = self.runner.sessions
        self.counter = 0
        with self.sessions.begin() as session:
            classroom = Classroom(title="ریاضی")
            session.add(classroom)
            session.flush()
            self.class_id = classroom.id
            student = Student(national_id="0012345678", full_name="دانش آموز", phone="09123456789")
            session.add(student)
            session.flush()
            self.student_id = student.id
            session.add(Enrollment(classroom_id=classroom.id, student_id=student.id))

    def tearDown(self):
        if self.runner.lock_connection is not None:
            self.runner.close()
            self.engine = create_engine(os.environ["TEST_DATABASE_URL"])
        Base.metadata.drop_all(self.engine)
        self.engine.dispose()

    def message(self, text="", user=25, document=None, chat_type="private", **extra):
        self.counter += 1
        message = {"message_id": self.counter, "from": {"id": user},
                   "chat": {"id": user, "type": chat_type}, "text": text, **extra}
        if document:
            message["document"] = document
        update = {"update_id": self.counter, "message": message}
        self.runner.ingest([update])
        self.runner.process_one()
        return update

    def callback(self, data, user=None):
        self.counter += 1
        user = self.config.admin_id if user is None else user
        self.runner.ingest([{"update_id": self.counter, "callback_query": {
            "id": f"2{self.counter}", "from": {"id": user}, "data": data,
            "message": {"chat": {"id": user, "type": "private"}},
        }}])
        self.runner.process_one()

    def latest_text(self):
        with self.sessions() as session:
            items = list(session.scalars(select(Outbox).where(Outbox.method == "sendMessage").order_by(Outbox.id)))
            return items[-1].payload["text"]

    def pending(self, kind="attendance", status="حاضر"):
        row = {"national_id": "0012345678", "name": "دانش آموز"}
        if kind == "attendance":
            row["status"] = status
        elif kind == "students":
            row["phone"] = "09123456789"
        else:
            row.update(code="OFF-20", description="دوره اول")
        return {"kind": kind, "classroom_id": self.class_id,
                "session_date": excel.today()[0].isoformat(), "rows": [row]}

    def test_login_persian_digits(self):
        self.message("۰۰۱۲۳۴۵۶۷۸")
        with self.sessions() as session:
            student = session.get(Student, self.student_id)
            self.assertEqual(student.chat_id, 25)
        self.assertIn("خوش آمدید", self.latest_text())

    def test_unknown_student_exact_error(self):
        self.message("9999999999")
        self.assertEqual(self.latest_text(), "کد ملی وارد شده در سامانه موجود نیست؛ لطفاً با مدیریت تماس بگیرید.")

    def test_invalid_national_id(self):
        self.message("123")
        self.assertIn("۱۰ رقم", self.latest_text())

    def test_cannot_take_over_account(self):
        self.message("0012345678")
        self.message("0012345678", user=26)
        self.assertIn("قبلاً متصل", self.latest_text())

    def test_login_rate_limit_survives_start(self):
        for _index in range(10):
            self.message("9999999999")
        self.message("/start")
        self.message("0012345678")
        self.assertIn("بیش از حد", self.latest_text())
        with self.sessions() as session:
            self.assertIsNone(session.get(Student, self.student_id).chat_id)

    def test_remove_enrollment_keeps_history(self):
        with self.sessions.begin() as session:
            apply_import(session, self.pending())
        self.callback(f"remove:{self.class_id}")
        self.message("0012345678", user=self.config.admin_id)
        with self.sessions() as session:
            nonce = session.get(Conversation, self.config.admin_id).data["nonce"]
        self.callback("confirm:" + nonce)
        with self.sessions() as session:
            self.assertIsNone(session.get(Enrollment, (self.class_id, self.student_id)))
            self.assertEqual(session.scalar(select(func.count()).select_from(Attendance)), 1)

    def test_unlink_allows_rebinding_with_confirmation(self):
        self.message("0012345678")
        self.callback(f"unlink:{self.student_id}")
        with self.sessions() as session:
            nonce = session.get(Conversation, self.config.admin_id).data["nonce"]
            self.assertEqual(session.get(Student, self.student_id).chat_id, 25)
        self.callback("confirm:" + nonce)
        self.message("0012345678", user=26)
        with self.sessions() as session:
            self.assertEqual(session.get(Student, self.student_id).chat_id, 26)

    def test_failed_update_has_backoff_and_does_not_skip_order(self):
        self.api.error = BaleError(503)
        self.callback(f"upload:attendance:{self.class_id}")
        from unittest.mock import patch

        with patch.object(self.api, "download", side_effect=BaleError(503)):
            self.message(user=self.config.admin_id, document={"file_id": "id", "file_name": "file.xlsx"})
        self.assertFalse(self.runner.process_one())
        with self.sessions() as session:
            failed = session.get(Inbox, self.counter)
            self.assertEqual(failed.status, "pending")
            self.assertEqual(failed.attempts, 1)

    def test_prune_expired_pending_import(self):
        self.message("➕ ایجاد کلاس", user=self.config.admin_id)
        with self.sessions.begin() as session:
            session.get(Conversation, self.config.admin_id).updated_at = utcnow() - timedelta(hours=3)
        self.runner.prune()
        with self.sessions() as session:
            self.assertEqual(session.get(Conversation, self.config.admin_id).data, {})

    def test_inactive_student_rejected(self):
        with self.sessions.begin() as session:
            session.get(Student, self.student_id).active = False
        self.message("0012345678")
        self.assertIn("موجود نیست", self.latest_text())

    def test_admin_menu_and_security(self):
        self.message("/start", user=self.config.admin_id)
        self.message("➕ ایجاد کلاس", user=25)
        self.callback(f"rename:{self.class_id}", user=25)
        with self.sessions() as session:
            self.assertEqual(session.get(Classroom, self.class_id).title, "ریاضی")
        self.assertIn("مجاز نیست", self.latest_text())

    def test_groups_are_ignored(self):
        self.message("/start", user=self.config.admin_id, chat_type="group")
        with self.sessions() as session:
            self.assertEqual(session.scalar(select(func.count()).select_from(Outbox)), 0)

    def test_create_rename_class(self):
        self.message("➕ ایجاد کلاس", user=self.config.admin_id)
        self.message("علوم", user=self.config.admin_id)
        self.callback(f"rename:{self.class_id}")
        self.message("ریاضی جدید", user=self.config.admin_id)
        with self.sessions() as session:
            self.assertEqual(session.get(Classroom, self.class_id).title, "ریاضی جدید")
            self.assertEqual(session.scalar(select(func.count()).select_from(Classroom)), 2)

    def test_contact_setting(self):
        self.message("⚙️ تنظیم راه‌های ارتباطی", user=self.config.admin_id)
        self.message("تماس: 02100000000", user=self.config.admin_id)
        self.message("0012345678")
        self.message("☎️ راه‌های ارتباطی")
        self.assertEqual(self.latest_text(), "تماس: 02100000000")

    def test_attendance_idempotent_and_changes_notify(self):
        self.message("0012345678")
        with self.sessions.begin() as session:
            self.assertEqual(apply_import(session, self.pending()), 1)
        with self.sessions.begin() as session:
            self.assertEqual(apply_import(session, self.pending()), 0)
            self.assertEqual(apply_import(session, self.pending(status="غایب")), 1)
        with self.sessions() as session:
            self.assertEqual(session.scalar(select(func.count()).select_from(Attendance)), 1)
            notices = [item for item in session.scalars(select(Outbox)) if "وضعیت شما" in item.payload.get("text", "")]
            self.assertEqual(len(notices), 2)
            self.assertIn(excel.today()[1], notices[0].payload["text"])

    def test_unbound_student_no_notice(self):
        with self.sessions.begin() as session:
            apply_import(session, self.pending())
        with self.sessions() as session:
            self.assertEqual(session.scalar(select(func.count()).select_from(Outbox)), 0)

    def test_in_bot_attendance_confirm_and_notify(self):
        self.message("0012345678")
        self.callback(f"attend:{self.class_id}")
        with self.sessions() as session:
            nonce = session.get(Conversation, self.config.admin_id).data["nonce"]
        self.callback(f"mark:{nonce}:{self.student_id}:absent")
        with self.sessions() as session:
            self.assertEqual(session.scalar(select(func.count()).select_from(Attendance)), 0)
        self.callback(f"confirm:{nonce}")
        with self.sessions() as session:
            self.assertEqual(session.scalar(select(Attendance)).status, "غایب")
            notices = [item for item in session.scalars(select(Outbox))
                       if item.payload.get("chat_id") == 25 and "وضعیت شما" in item.payload.get("text", "")]
            self.assertEqual(len(notices), 1)
        self.message("📋 حضور و غیاب من")
        self.assertIn("غایب", self.latest_text())
        self.assertIn("ریاضی", self.latest_text())

    def test_cancel_in_bot_attendance_rejects_old_buttons(self):
        self.callback(f"attend:{self.class_id}")
        with self.sessions() as session:
            nonce = session.get(Conversation, self.config.admin_id).data["nonce"]
        self.message("لغو", user=self.config.admin_id)
        self.callback(f"mark:{nonce}:{self.student_id}:present")
        self.assertIn("منقضی", self.latest_text())
        with self.sessions() as session:
            self.assertEqual(session.scalar(select(func.count()).select_from(Attendance)), 0)

    def test_student_attendance_history_is_private(self):
        with self.sessions.begin() as session:
            apply_import(session, self.pending())
        self.message("📋 حضور و غیاب من", user=26)
        self.assertNotIn("ریاضی", self.latest_text())

    def test_import_rejects_foreign_student_atomically(self):
        pending = self.pending()
        pending["rows"].append({"national_id": "9999999999", "name": "ناشناس", "status": "حاضر"})
        with self.assertRaises(ValueError):
            with self.sessions.begin() as session:
                apply_import(session, pending)
        with self.sessions() as session:
            self.assertEqual(session.scalar(select(func.count()).select_from(Attendance)), 0)

    def test_student_import_preserves_binding_and_membership(self):
        self.message("0012345678")
        with self.sessions.begin() as session:
            pending = self.pending("students")
            pending["rows"][0]["name"] = "نام جدید"
            apply_import(session, pending)
        with self.sessions() as session:
            student = session.get(Student, self.student_id)
            self.assertEqual(student.full_name, "نام جدید")
            self.assertEqual(student.chat_id, 25)
            self.assertEqual(session.scalar(select(func.count()).select_from(Enrollment)), 1)

    def test_discount_upsert_and_personal_menu(self):
        self.message("0012345678")
        with self.sessions.begin() as session:
            apply_import(session, self.pending("discounts"))
        with self.sessions.begin() as session:
            pending = self.pending("discounts")
            pending["rows"][0]["description"] = "دوره جدید"
            apply_import(session, pending)
        self.message("🎁 تخفیف‌های من")
        self.assertIn("OFF-20", self.latest_text())
        self.assertIn("دوره جدید", self.latest_text())
        with self.sessions() as session:
            self.assertEqual(session.scalar(select(func.count()).select_from(Discount)), 1)

    def test_upload_requires_confirmation_and_is_single_use(self):
        self.api.content = excel.workbook("attendance", [("0012345678", "دانش آموز", "حاضر")], self.class_id,
                                          excel.today()[0].isoformat())
        self.callback(f"upload:attendance:{self.class_id}")
        self.message(user=self.config.admin_id, document={"file_id": "test", "file_name": "test.xlsx"})
        with self.sessions() as session:
            self.assertEqual(session.scalar(select(func.count()).select_from(Attendance)), 0)
            nonce = session.get(Conversation, self.config.admin_id).data["nonce"]
        self.callback("confirm:" + nonce)
        self.callback("confirm:" + nonce)
        with self.sessions() as session:
            self.assertEqual(session.scalar(select(func.count()).select_from(Attendance)), 1)
        self.assertIn("قبلاً انجام شده", self.latest_text())

    def test_cancel_upload(self):
        self.callback(f"upload:students:{self.class_id}")
        self.message("لغو", user=self.config.admin_id)
        with self.sessions() as session:
            self.assertEqual(session.get(Conversation, self.config.admin_id).data, {})

    def test_broadcast_media_confirm_deduplicated(self):
        self.message("0012345678")
        self.message("📢 پیام همگانی", user=self.config.admin_id)
        self.message(user=self.config.admin_id, photo=[{"file_id": "photo"}])
        with self.sessions() as session:
            self.assertEqual(session.scalar(select(func.count()).select_from(Outbox).where(Outbox.method == "copyMessage")), 0)
            nonce = session.get(Conversation, self.config.admin_id).data["nonce"]
        self.callback("confirm:" + nonce)
        self.callback("confirm:" + nonce)
        with self.sessions() as session:
            copies = list(session.scalars(select(Outbox).where(Outbox.method == "copyMessage")))
            self.assertEqual(len(copies), 1)
            self.assertEqual(copies[0].payload["chat_id"], 25)

    def test_class_broadcast_excludes_other_class(self):
        self.message("0012345678")
        with self.sessions.begin() as session:
            classroom = Classroom(title="دیگر")
            session.add(classroom)
            session.flush()
            self.assertEqual(self.runner.bot.recipients(session, classroom.id), [])
            self.assertEqual(self.runner.bot.recipients(session, self.class_id), [25])

    def test_inbox_duplicate_not_processed_twice(self):
        update = self.message("/start", user=self.config.admin_id)
        self.runner.ingest([update])
        self.assertFalse(self.runner.process_one())
        self.assertEqual(self.runner.offset(), update["update_id"] + 1)
        with self.sessions() as session:
            self.assertEqual(session.scalar(select(func.count()).select_from(Inbox)), 1)

    def test_delivery_retry_after_and_failure(self):
        self.message("/start", user=self.config.admin_id)
        self.api.error = BaleError(429, 30)
        self.runner.deliver_one()
        with self.sessions.begin() as session:
            item = session.scalar(select(Outbox))
            self.assertEqual(item.status, "pending")
            self.assertEqual(item.attempts, 1)
            item.available_at = utcnow() - timedelta(seconds=1)
        self.api.error = BaleError(403)
        self.runner.deliver_one()
        with self.sessions() as session:
            self.assertEqual(session.scalar(select(Outbox)).status, "failed")

    def test_delivery_success_scrubs_payload(self):
        self.message("/start", user=self.config.admin_id)
        self.runner.deliver_one()
        with self.sessions() as session:
            item = session.scalar(select(Outbox))
            self.assertEqual(item.status, "sent")
            self.assertEqual(item.payload, {})

    def test_state_survives_runner_recreation(self):
        self.message("➕ ایجاد کلاس", user=self.config.admin_id)
        second = Runner(self.config, self.api, self.engine)
        with second.sessions() as session:
            self.assertEqual(session.get(Conversation, self.config.admin_id).data["mode"], "create")

    def test_expired_confirmation_rejected(self):
        self.message("📢 پیام همگانی", user=self.config.admin_id)
        self.message("خبر", user=self.config.admin_id)
        with self.sessions.begin() as session:
            state = session.get(Conversation, self.config.admin_id)
            nonce = state.data["nonce"]
            state.updated_at = utcnow() - timedelta(hours=3)
        self.callback("confirm:" + nonce)
        self.assertIn("منقضی", self.latest_text())


class ExcelTests(unittest.TestCase):
    def file(self, rows, kind="students", classroom_id=7):
        return excel.workbook(kind, rows, classroom_id, excel.today()[0].isoformat())

    def test_round_trip_leading_zero_and_metadata(self):
        content = self.file([("نام", "0012345678", "09123456789")])
        parsed = excel.parse(content, "students", 7)
        self.assertEqual(parsed["rows"][0]["national_id"], "0012345678")
        self.assertEqual(parsed["rows"][0]["phone"], "09123456789")

    def test_wrong_class_and_kind(self):
        content = self.file([("نام", "0012345678", "09123456789")])
        for kind, classroom in [("students", 8), ("attendance", 7)]:
            with self.assertRaises(excel.ExcelError):
                excel.parse(content, kind, classroom)

    def test_duplicate_row_rejected(self):
        row = ("نام", "0012345678", "09123456789")
        with self.assertRaises(excel.ExcelError):
            excel.parse(self.file([row, row]), "students", 7)

    def test_bad_status_and_blank_rejected(self):
        for status in ("", "تاخیر", "present"):
            with self.assertRaises(excel.ExcelError):
                excel.parse(self.file([("0012345678", "نام", status)], "attendance"), "attendance", 7)

    def test_formula_injection_is_text_on_export(self):
        content = self.file([("=HYPERLINK(1)", "0012345678", "09123456789")])
        book = load_workbook(BytesIO(content))
        self.assertEqual(book["داده‌ها"]["A2"].data_type, "s")
        book.close()

    def test_actual_formula_on_upload_rejected(self):
        content = self.file([("نام", "0012345678", "09123456789")])
        book = load_workbook(BytesIO(content))
        book["داده‌ها"]["A2"] = "=1+1"
        output = BytesIO()
        book.save(output)
        book.close()
        with self.assertRaises(excel.ExcelError):
            excel.parse(output.getvalue(), "students", 7)

    def test_max_rows(self):
        content = self.file([("نام", "0012345678", "09123456789"), ("نام دوم", "0012345679", "09123456789")])
        with self.assertRaises(excel.ExcelError):
            excel.parse(content, "students", 7, max_rows=1)

    def test_bad_archive_and_empty_file(self):
        for content in (b"not excel", self.file([])):
            with self.assertRaises(excel.ExcelError):
                excel.parse(content, "students", 7)

    def test_jalali_tehran_today(self):
        current_date, jalali = excel.today()
        import jdatetime

        self.assertEqual(jdatetime.datetime.strptime(jalali, "%Y/%m/%d").date().togregorian(), current_date)


if __name__ == "__main__":
    unittest.main()
